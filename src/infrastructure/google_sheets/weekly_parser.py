import re

from openpyxl.worksheet.worksheet import Worksheet

from src.domain.constants import PAIR_TIMES
from src.domain.models import ScheduleLesson
from src.infrastructure.google_sheets.parsing import is_class_header, parse_integer, parse_time_range, split_subject_teacher, text, weekday_number


def parse_weekly_sheet(sheet: Worksheet) -> tuple[list[ScheduleLesson], list[str]]:
    class_columns = {
        column: text(sheet.cell(1, column).value)
        for column in range(1, sheet.max_column + 1)
        if is_class_header(sheet.cell(1, column).value)
    }
    if sheet.title == "Начальная школа":
        class_columns = {column: value for column, value in class_columns.items() if column in (4, 7, 10, 13)}
    available_classes = sorted(set(class_columns.values()), key=str.casefold)
    if not class_columns:
        return [], available_classes

    active_days: dict[int, int] = {}
    active_pair_number: int | None = None
    lessons: list[ScheduleLesson] = []
    for row in range(1, sheet.max_row + 1):
        pair_cell = sheet.cell(row, 3).value
        if sheet.title == "Начальная школа":
            pair_cell = sheet.cell(row, 2).value
        parsed_pair = parse_integer(pair_cell)
        if parsed_pair is not None:
            active_pair_number = parsed_pair
        day_header = False
        for column in class_columns:
            weekday = weekday_number(sheet.cell(row, column).value)
            if weekday is not None:
                active_days[column] = weekday
                day_header = True
        if day_header:
            continue
        for column, class_course in class_columns.items():
            weekday = active_days.get(column)
            if weekday is None:
                continue
            subject_raw = text(sheet.cell(row, column).value)
            if not subject_raw or weekday_number(subject_raw) is not None:
                continue
            if "время уроков" in subject_raw.casefold() or "куратор" in subject_raw.casefold() or "преподаватели классики" in subject_raw.casefold():
                continue
            lesson_number = parse_integer(sheet.cell(row, 1).value)
            pair_number = active_pair_number
            if sheet.title == "Начальная школа":
                lesson_number = parse_integer(sheet.cell(row, column - 1).value)
                pair_number = active_pair_number
                lesson_raw = sheet.cell(row, 1).value
                pair_raw = sheet.cell(row, 2).value
                lesson_time = parse_time_range(lesson_raw)
                pair_time = parse_time_range(pair_raw)
                raw_time = text(lesson_raw) or text(pair_raw) or None
            else:
                lesson_raw = sheet.cell(row, 2).value
                pair_raw = sheet.cell(row, 4).value
                lesson_time = parse_time_range(lesson_raw)
                pair_time = parse_time_range(pair_raw)
                raw_time = text(lesson_raw) or text(pair_raw) or None
            if pair_number in PAIR_TIMES:
                pair_time = PAIR_TIMES[pair_number]
            room_value = text(sheet.cell(row, column + 1).value)
            entries = _split_group_entries(subject_raw)
            for entry_raw, group in entries:
                subject, teacher, _ = split_subject_teacher(entry_raw)
                lessons.append(ScheduleLesson(
                    weekday=weekday,
                    class_course=class_course,
                    group=group,
                    lesson_number=lesson_number,
                    pair_number=pair_number,
                    start_time=lesson_time[0],
                    end_time=lesson_time[1],
                    pair_start_time=pair_time[0],
                    pair_end_time=pair_time[1],
                    subject=subject,
                    subject_raw=entry_raw,
                    teacher=teacher,
                    room=room_value or None,
                    source_sheet=sheet.title,
                    source_row=row,
                    time_raw=raw_time,
                ))
    return lessons, available_classes


def _split_group_entries(raw: str) -> list[tuple[str, str | None]]:
    if "//" not in raw:
        match = re.search(r"\b(Д\s*[12]?|М)\b/?\s*$", raw, re.IGNORECASE)
        if match:
            return [(raw[:match.start()].strip(), re.sub(r"\s+", "", match.group(1)).upper())]
        return [(raw, None)]
    entries: list[tuple[str, str | None]] = []
    for part in raw.split("//"):
        match = re.search(r"\b(Д\s*[12]?|М)\b/?\s*$", part, re.IGNORECASE)
        group = re.sub(r"\s+", "", match.group(1)).upper() if match else None
        entries.append((part[:match.start()].strip() if match else part.strip(), group))
    return entries
