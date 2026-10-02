from datetime import date, datetime

from openpyxl.worksheet.worksheet import Worksheet

from src.infrastructure.google_sheets.parsing import parse_date, text


def parse_rooms(sheet: Worksheet) -> dict[date, list[tuple[str, str, int | None, str | None]]]:
    result: dict[date, list[tuple[str, str, int | None, str | None]]] = {}
    current_date: date | None = None
    room_headers: dict[int, str] = {}
    for row in range(1, sheet.max_row + 1):
        date_value = sheet.cell(row, 3).value
        parsed_date = parse_date(date_value)
        if parsed_date:
            current_date = parsed_date
            result.setdefault(current_date, [])
            room_headers = {column: text(sheet.cell(row + 1, column).value) for column in range(3, sheet.max_column + 1)}
            continue
        if not current_date:
            continue
        slot_text = text(sheet.cell(row, 2).value)
        if "пара" not in slot_text.casefold():
            continue
        import re
        number_match = re.search(r"(\d+)\s*пара", slot_text.casefold())
        pair_number = int(number_match.group(1)) if number_match else None
        for column in range(3, sheet.max_column + 1):
            value = text(sheet.cell(row, column).value)
            if not value:
                continue
            room = room_headers.get(column, "")
            result[current_date].append((room, value, pair_number, slot_text))
    return result
