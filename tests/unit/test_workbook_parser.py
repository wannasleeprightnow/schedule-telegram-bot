from datetime import date, time

from openpyxl import load_workbook

from src.domain.constants import PAIR_TIMES
from src.domain.services import build_schedule
from src.infrastructure.google_sheets.changes_parser import parse_changes
from src.infrastructure.google_sheets.parsing import parse_date, parse_time_range, split_subject_teacher
from src.infrastructure.google_sheets.rooms_parser import parse_rooms
from src.infrastructure.google_sheets.workbook_parser import parse_workbook


def test_parses_supported_sheets_and_discovers_classes(workbook_path):
    result = parse_workbook(str(workbook_path))
    assert len(result.lessons) > 500
    assert result.available_classes["primary"]
    assert "III (8) курс" in result.available_classes["course"]
    assert "5 (9) Б" in result.available_classes["school"]
    assert all(lesson.source_sheet not in {"Информация (старое)", "Лист37", "Лист38"} for lesson in result.lessons)


def test_time_teacher_and_invalid_time_are_safe():
    assert parse_time_range("09:00–09:45") == (time(9), time(9, 45))
    assert parse_time_range("15:60 - 16:15") == (None, None)
    assert parse_time_range("not a time") == (None, None)
    assert parse_time_range("15:00:00") == (time(15), None)
    subject, teacher, _ = split_subject_teacher("История (Делигиоз Г.Г.)")
    assert subject == "История"
    assert teacher == "Делигиоз Г.Г."


def test_weekly_parser_extracts_dates_times_teachers_and_pairs(workbook_path):
    result = parse_workbook(str(workbook_path))
    lesson = next(item for item in result.lessons if item.class_course == "III (8) курс" and item.weekday == 4 and item.subject.startswith("Дуэтно-классический"))
    assert lesson.pair_number == 3
    assert lesson.pair_start_time == time(13)
    assert lesson.pair_end_time == time(14, 30)
    assert lesson.teacher == "Матвеев К.С."


def test_all_pair_times_follow_the_academy_bell_schedule(workbook_path):
    result = parse_workbook(str(workbook_path))
    paired_lessons = [lesson for lesson in result.lessons if lesson.pair_number in PAIR_TIMES]
    assert paired_lessons
    assert all(
        (lesson.pair_start_time, lesson.pair_end_time) == PAIR_TIMES[lesson.pair_number]
        for lesson in paired_lessons
    )


def test_information_date_forward_fill_and_public_announcements(workbook_path):
    workbook = load_workbook(workbook_path, data_only=True)
    changes, announcements = parse_changes(workbook["ИНФОРМАЦИЯ 2026"])
    assert any(item.date == date(2026, 9, 2) and item.class_course == "3 (7) Б" for item in changes)
    assert any(item.date == date(2026, 9, 2) and item.class_course == "I (6) курс" for item in changes)
    assert any(item.date == date(2026, 9, 11) and "Охрана труда" in item.information for item in changes)
    assert any(item.date == date(2026, 9, 7) and "плаванию" in item.information for item in announcements)
    assert not any("смартфонов" in item.information for item in announcements)


def test_russian_date_formats_and_latin_c_typo():
    assert parse_date("02 cентября (среда)") == date(2026, 9, 2)
    assert parse_date("2.10.2026") == date(2026, 10, 2)


def test_rooms_sheet_has_real_dated_room_assignments(workbook_path):
    workbook = load_workbook(workbook_path, data_only=True)
    rooms = parse_rooms(workbook["Залы"])
    entries = rooms[date(2026, 10, 2)]
    assert any(pair == 1 and "5 (9) Б" in cell and room == "Зал № 2" for room, cell, pair, _ in entries)


def test_real_date_regressions(workbook_path):
    data = parse_workbook(str(workbook_path))
    target_dates = [date(2026, 9, 2), date(2026, 9, 11), date(2026, 9, 12), date(2026, 9, 23), date(2026, 9, 28), date(2026, 10, 2)]
    assert all(any(change.date == requested for change in data.changes + data.announcements) for requested in target_dates)

    canceled = build_schedule(date(2026, 10, 2), "III (8) курс", None, data.lessons, data.changes, data.announcements)
    assert not any("охрана труда артиста балета" in lesson.subject.casefold() for lesson in canceled.lessons)
    assert not any("охрана труда артиста балета" in item.casefold() for item in canceled.unresolved_changes)
    psychology = next(lesson for lesson in canceled.lessons if lesson.pair_number == 3)
    assert psychology.subject == "Психология общения"
    assert psychology.teacher == "Леушина А.В."
    assert psychology.room == "3.04"
    assert psychology.is_replacement
    assert "Психология общения" not in canceled.unresolved_changes

    september_canceled = build_schedule(date(2026, 9, 2), "3 (7) Б", None, data.lessons, data.changes, data.announcements)
    assert not any("исполнительская практика" in lesson.subject.casefold() for lesson in september_canceled.lessons)

    september_replacement = build_schedule(date(2026, 9, 11), "III (8) курс", None, data.lessons, data.changes, data.announcements)
    assert not any("дуэтно-классический танец" in lesson.subject.casefold() for lesson in september_replacement.lessons)
    replacement = next(lesson for lesson in september_replacement.lessons if lesson.pair_number == 3)
    assert replacement.subject == "Охрана труда"
    assert replacement.teacher == "Сорокина М.В."
    assert replacement.room == "4.06"
    assert replacement.is_replacement

    joined = build_schedule(date(2026, 10, 2), "5 (9) Б", None, data.lessons, data.changes, data.announcements)
    assert any("объединенный урок" in lesson.subject.casefold() for lesson in joined.lessons)

    september_joined = build_schedule(date(2026, 9, 12), "4 (8) А", None, data.lessons, data.changes, data.announcements)
    assert any("объединенный урок" in lesson.subject.casefold() for lesson in september_joined.lessons)

    added = build_schedule(date(2026, 9, 28), "4 (8) А", None, data.lessons, data.changes, data.announcements)
    exam = next(lesson for lesson in added.lessons if "экзамен по дисциплине" in lesson.subject.casefold())
    assert exam.start_time == time(17)
    assert exam.end_time == time(17, 30)
